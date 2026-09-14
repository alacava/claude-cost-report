#!/usr/bin/env python3
"""
claude_team_cost_report.py — per-user cost report for a Claude Team org.

Accepts EITHER export from claude.ai analytics (format auto-detected):

  1. Spend report        (Analytics > "How much is Claude costing?" >
                          Export spend report; supports custom date ranges
                          up to 90 days back). One row per user/product/model.
  2. Members analytics   (members-analytics-*.csv). One row per member,
                          rolled-up "Estimated Spend (USD)".

Joins in your seat fees (which no export contains) and writes a console
summary + an .xlsx with live formulas. With a spend report, a
"By Product & Model" detail sheet is added, and you can pass the members
export as --roster so zero-usage seat holders still appear with their
seat cost.

Optionally, --console-usage pulls platform.claude.com (Claude Console /
Claude Platform API) cost data via the Admin API and adds it as its own
"Console API Usage" sheet. This is a SEPARATE billing relationship from
Team seats (pay-as-you-go API usage, not tied to named Team members), so
it is never blended into Per-User Cost — see CLAUDE.md.

Usage:
    python3 claude_team_cost_report.py <export.csv> [options]

Options:
    --roster FILE       members-analytics CSV to supply the full member
                        list and Seat Tier (recommended with a spend report)
    --seat-price N      Price per Standard seat for the period (default 30.00,
                        Team Standard monthly; use 25 for annual)
    --premium-price N   Price per Premium seat (default 150.00 monthly)
    --out FILE          Output .xlsx path (default: derived from input name)
    --no-xlsx           Console summary only
    --console-usage     Also fetch platform.claude.com cost data via the
                        Admin API (requires ANTHROPIC_ADMIN_KEY env var)
    --console-start DATE  Start date (YYYY-MM-DD) for Console/API data;
                        default: same period as the Team CSV filename
    --console-end DATE  End date (YYYY-MM-DD) for Console/API data;
                        default: same period as the Team CSV filename

Requires: pandas, openpyxl
"""

import argparse
import json
import os
import re
import sys
import urllib.error
import urllib.parse
import urllib.request
from datetime import date, timedelta
from pathlib import Path

import pandas as pd

ANTHROPIC_API_BASE = "https://api.anthropic.com"
ANTHROPIC_VERSION = "2023-06-01"

MEMBERS_SPEND = "Estimated Spend (USD)"
SPEND_NET = "total_net_spend_usd"
SPEND_GROSS = "total_gross_spend_usd"


def detect_format(df) -> str:
    if {"user_email", SPEND_NET, "product", "model"} <= set(df.columns):
        return "spend"
    if {"Email", "Seat Tier", MEMBERS_SPEND} <= set(df.columns):
        return "members"
    sys.exit("Unrecognized CSV. Expected a claude.ai spend report "
             "(user_email/total_net_spend_usd/...) or members-analytics "
             "export (Email/Seat Tier/Estimated Spend (USD)).")


def seat_fee(tier: str, std: float, prem: float) -> float:
    t = (tier or "").strip().lower()
    if t == "standard":
        return std
    if t == "premium":
        return prem
    return 0.0  # Unassigned / unknown tiers carry no seat fee


def parse_period_dates(path: Path):
    """Return (start, end) YYYY-MM-DD strings from an export filename, or
    (None, None) if the filename doesn't carry a date range."""
    m = re.search(r"(\d{4}-\d{2}-\d{2})-to-(\d{4}-\d{2}-\d{2})", path.name)
    return (m.group(1), m.group(2)) if m else (None, None)


def period_from_filename(path: Path) -> str:
    start, end = parse_period_dates(path)
    return f"{start} to {end}" if start else "unknown (not in filename)"


def _console_api_get(admin_key: str, path: str, params) -> dict:
    """GET against the Console Admin API (platform.claude.com), stdlib-only."""
    url = f"{ANTHROPIC_API_BASE}{path}?{urllib.parse.urlencode(params, doseq=True)}"
    req = urllib.request.Request(url, headers={
        "x-api-key": admin_key,
        "anthropic-version": ANTHROPIC_VERSION,
        "User-Agent": "claude-cost-report/1.0",
    })
    try:
        with urllib.request.urlopen(req, timeout=30) as resp:
            return json.load(resp)
    except urllib.error.HTTPError as e:
        body = e.read().decode("utf-8", "replace")
        sys.exit(f"Console/API request to {path} failed: HTTP {e.code}\n{body}")
    except urllib.error.URLError as e:
        sys.exit(f"Console/API request to {path} failed: {e.reason}")


def fetch_console_cost_report(admin_key: str, start_date: str, end_date: str):
    """Pull the Admin API cost report for [start_date, end_date] (inclusive
    calendar days). Returns a flat list of per-item cost dicts across all
    daily buckets and pages."""
    starting_at = f"{start_date}T00:00:00Z"
    ending_at_date = date.fromisoformat(end_date) + timedelta(days=1)
    ending_at = f"{ending_at_date.isoformat()}T00:00:00Z"

    rows = []
    page = None
    while True:
        params = [("starting_at", starting_at), ("ending_at", ending_at),
                  ("group_by[]", "workspace_id"), ("group_by[]", "description"),
                  ("limit", 31)]
        if page:
            params.append(("page", page))
        data = _console_api_get(admin_key, "/v1/organizations/cost_report", params)
        for bucket in data.get("data", []):
            for item in bucket.get("results", []):
                rows.append({
                    "workspace_id": item.get("workspace_id"),
                    "model": item.get("model"),
                    "cost_type": item.get("cost_type"),
                    "service_tier": item.get("service_tier"),
                    # amount is a decimal string in the currency's lowest
                    # unit (cents for USD): "123.45" == $1.23
                    "amount_usd": float(item.get("amount", 0)) / 100.0,
                })
        page = data.get("next_page")
        if not data.get("has_more") or not page:
            break
    return rows


def fetch_console_workspace_names(admin_key: str) -> dict:
    """Map workspace_id -> display name via the Admin API (best-effort)."""
    names = {}
    after = None
    while True:
        params = [("limit", 1000), ("include_archived", "true")]
        if after:
            params.append(("after_id", after))
        data = _console_api_get(admin_key, "/v1/organizations/workspaces", params)
        for ws in data.get("data", []):
            names[ws["id"]] = ws.get("name") or ws["id"]
        after = data.get("last_id")
        if not data.get("has_more") or not after:
            break
    return names


def build_console_df(rows, workspace_names: dict) -> pd.DataFrame:
    """Aggregate raw cost-report line items into one row per
    workspace/model/cost_type/service_tier, sorted by spend descending."""
    cols = ["workspace", "model", "cost_type", "service_tier", "amount_usd"]
    if not rows:
        return pd.DataFrame(columns=cols)
    df = pd.DataFrame(rows)
    df["workspace"] = df["workspace_id"].apply(
        lambda w: workspace_names.get(w, w) if pd.notna(w) else "Default workspace")
    df["model"] = df["model"].fillna("—")
    df["cost_type"] = df["cost_type"].fillna("—")
    df["service_tier"] = df["service_tier"].fillna("—")
    return (df.groupby(["workspace", "model", "cost_type", "service_tier"],
                       as_index=False)
              .agg(amount_usd=("amount_usd", "sum"))
              .sort_values("amount_usd", ascending=False)
              .reset_index(drop=True))[cols]


def load_roster(path: Path):
    r = pd.read_csv(path)
    need = {"Email", "Seat Tier"}
    if not need <= set(r.columns):
        sys.exit(f"--roster file is missing {need - set(r.columns)}; "
                 "expected the members-analytics export.")
    r = r.rename(columns={"Email": "email", "Seat Tier": "seat_tier",
                          "Name": "name", "Role": "role"})
    keep = [c for c in ("email", "name", "role", "seat_tier") if c in r.columns]
    r["email"] = r["email"].str.strip().str.lower()
    return r[keep].drop_duplicates("email")


def normalize(df, fmt, roster):
    """Return (per_user_df, detail_df_or_None).

    per_user_df columns: email, name, role, seat_tier, usage_spend
    (+ gross_spend/requests for the spend format)."""
    detail = None
    if fmt == "members":
        users = df.rename(columns={
            "Email": "email", "Name": "name", "Role": "role",
            "Seat Tier": "seat_tier", MEMBERS_SPEND: "usage_spend"})
        users["email"] = users["email"].str.strip().str.lower()
        users = users[["email", "name", "role", "seat_tier", "usage_spend"]]
    else:
        d = df.copy()
        d["user_email"] = d["user_email"].str.strip().str.lower()
        for c in (SPEND_NET, SPEND_GROSS):
            d[c] = pd.to_numeric(d[c], errors="coerce").fillna(0.0)
        detail = (d.groupby(["user_email", "product", "model"], as_index=False)
                    .agg(requests=("total_requests", "sum"),
                         prompt_tokens=("total_prompt_tokens", "sum"),
                         completion_tokens=("total_completion_tokens", "sum"),
                         net_spend=(SPEND_NET, "sum"),
                         gross_spend=(SPEND_GROSS, "sum"))
                    .sort_values("net_spend", ascending=False))
        users = (d.groupby("user_email", as_index=False)
                   .agg(usage_spend=(SPEND_NET, "sum"),
                        gross_spend=(SPEND_GROSS, "sum"),
                        requests=("total_requests", "sum"))
                   .rename(columns={"user_email": "email"}))
        users["name"] = users["email"].str.split("@").str[0]
        users["role"] = ""
        users["seat_tier"] = "Standard"  # overridden by roster when given

    if roster is not None:
        users = roster.merge(users.drop(columns=[c for c in ("name", "role",
                                                             "seat_tier")
                                                 if c in users.columns]),
                             on="email", how="outer")
        for c in ("usage_spend", "gross_spend", "requests"):
            if c in users.columns:
                users[c] = users[c].fillna(0.0)
        users["seat_tier"] = users["seat_tier"].fillna("Standard")
        users["name"] = users["name"].fillna(users["email"].str.split("@").str[0])
    users["usage_spend"] = pd.to_numeric(users["usage_spend"],
                                         errors="coerce").fillna(0.0)
    return users, detail


def build_xlsx(users, detail, out_path, std_price, prem_price, period, fmt,
               roster_used, console_df=None, console_period=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    A = "Arial"
    base = Font(name=A, size=10)
    bold = Font(name=A, size=10, bold=True)
    blue_in = Font(name=A, size=10, color="0000FF")
    green_link = Font(name=A, size=10, color="008000")
    hdr_font = Font(name=A, size=10, bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="1F3864")
    yellow = PatternFill("solid", fgColor="FFFF00")
    gray = PatternFill("solid", fgColor="F2F2F2")
    thin = Border(bottom=Side(style="thin", color="BFBFBF"))
    top = Border(top=Side(style="thin", color="000000"))
    money = "$#,##0.00;($#,##0.00);-"

    ws = wb.active
    ws.title = "Assumptions"
    ws["A1"] = "Claude Team Per-User Cost Report — Assumptions"
    ws["A1"].font = Font(name=A, bold=True, size=13)
    src = ("spend report (per-user/per-model net spend)" if fmt == "spend"
           else "members-analytics export (rolled-up estimated spend)")
    rows = [
        ("Report period", period),
        ("Source export", src),
        ("Standard seat price ($/member/period)", std_price),
        ("Premium seat price ($/member/period)", prem_price),
        ("Unassigned / no-seat price", 0),
    ]
    r = 3
    for label, val in rows:
        ws.cell(r, 1, label).font = bold
        c = ws.cell(r, 2, val)
        if isinstance(val, (int, float)):
            c.font, c.fill = blue_in, yellow
        else:
            c.font = base
        r += 1
    notes = [
        "Yellow cells with blue text are inputs — edit to match your billing.",
        "Team Standard: $30/member/mo monthly or $25/member/mo annual (see support.claude.com).",
        "Seat fees are user-entered assumptions; the exports contain usage spend only.",
        "Usage spend = net spend after discounts/credits, straight from the export.",
        "Per-user total = seat fee + usage spend for the period.",
        f"Example: a Standard seat with $0 usage spend costs ${std_price:,.2f} at the current input.",
    ]
    if fmt == "spend" and not roster_used:
        notes.append("Seat tiers assumed Standard for every user in the spend "
                     "report — rerun with --roster <members-analytics.csv> for "
                     "actual tiers and zero-usage members.")
    if console_df is not None:
        notes.append(
            f"Console/API (platform.claude.com) spend for {console_period}: "
            f"${console_df['amount_usd'].sum():,.2f} — a separate pay-as-you-go "
            "billing relationship from Team seats, not tied to named members. "
            "Shown on its own 'Console API Usage' sheet, NOT included in the "
            "Per-User Cost TOTAL above.")
    ws.cell(9, 1, "Legend / how to use").font = bold
    for i, n in enumerate(notes):
        ws.cell(10 + i, 1, n).font = base
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 60

    ws2 = wb.create_sheet("Per-User Cost")
    has_req = "requests" in users.columns
    headers = ["Name", "Email", "Role", "Seat Tier"]
    if has_req:
        headers.append("Requests")
    headers += ["Usage Spend (USD)", "Seat Fee (USD)", "Total Cost (USD)",
                "% of Total"]
    for col, h in enumerate(headers, 1):
        c = ws2.cell(1, col, h)
        c.font, c.fill = hdr_font, hdr_fill
        c.alignment = Alignment(horizontal="center", vertical="center",
                                wrap_text=True)
    ws2.freeze_panes = "A2"

    n = len(users)
    first, last = 2, 1 + n
    spend_col = 5 + (1 if has_req else 0)
    fee_col, total_col, pct_col = spend_col + 1, spend_col + 2, spend_col + 3
    SL, FL, TL = (get_column_letter(x) for x in (spend_col, fee_col, total_col))

    for i, row in users.reset_index(drop=True).iterrows():
        r = 2 + i
        vals = [row.get("name", ""), row["email"], row.get("role", ""),
                row["seat_tier"]]
        if has_req:
            vals.append(int(row.get("requests", 0)))
        vals.append(row["usage_spend"])
        for col, v in enumerate(vals, 1):
            c = ws2.cell(r, col, v)
            c.font, c.border = base, thin
            if i % 2:
                c.fill = gray
        fee = ws2.cell(
            r, fee_col,
            f'=IF($D{r}="Standard",Assumptions!$B$5,'
            f'IF($D{r}="Premium",Assumptions!$B$6,Assumptions!$B$7))')
        fee.font = green_link
        ws2.cell(r, total_col, f"={SL}{r}+{FL}{r}").font = base
        ws2.cell(r, pct_col,
                 f"=IF(${TL}${last + 1}=0,0,{TL}{r}/${TL}${last + 1})").font = base
        for col in (fee_col, total_col, pct_col):
            ws2.cell(r, col).border = thin
            if i % 2:
                ws2.cell(r, col).fill = gray

    tr = last + 1
    ws2.cell(tr, 1, "TOTAL").font = bold
    sum_cols = ([5] if has_req else []) + [spend_col, fee_col, total_col]
    for col in sum_cols:
        L = get_column_letter(col)
        ws2.cell(tr, col, f"=SUM({L}{first}:{L}{last})").font = bold
    PL = get_column_letter(pct_col)
    ws2.cell(tr, pct_col,
             f"=IF({TL}{tr}=0,0,SUM({PL}{first}:{PL}{last}))").font = bold
    for col in range(1, pct_col + 1):
        ws2.cell(tr, col).border = top

    for r in range(first, tr + 1):
        for col in (spend_col, fee_col, total_col):
            ws2.cell(r, col).number_format = money
        ws2.cell(r, pct_col).number_format = "0.0%"
        if has_req:
            ws2.cell(r, 5).number_format = "#,##0"
    widths = [18, 34, 14, 12] + ([10] if has_req else []) + [16, 13, 14, 10]
    for i, w in enumerate(widths, 1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    if detail is not None and len(detail):
        ws3 = wb.create_sheet("By Product & Model")
        d_headers = ["Email", "Product", "Model", "Requests", "Prompt Tokens",
                     "Completion Tokens", "Net Spend (USD)", "Gross Spend (USD)"]
        for col, h in enumerate(d_headers, 1):
            c = ws3.cell(1, col, h)
            c.font, c.fill = hdr_font, hdr_fill
            c.alignment = Alignment(horizontal="center", vertical="center",
                                    wrap_text=True)
        ws3.freeze_panes = "A2"
        for i, row in detail.reset_index(drop=True).iterrows():
            r = 2 + i
            vals = [row["user_email"], row["product"], row["model"],
                    int(row["requests"]), int(row["prompt_tokens"]),
                    int(row["completion_tokens"]), row["net_spend"],
                    row["gross_spend"]]
            for col, v in enumerate(vals, 1):
                c = ws3.cell(r, col, v)
                c.font, c.border = base, thin
                if i % 2:
                    c.fill = gray
        dtr = 2 + len(detail)
        ws3.cell(dtr, 1, "TOTAL").font = bold
        for col in range(4, 9):
            L = get_column_letter(col)
            ws3.cell(dtr, col, f"=SUM({L}2:{L}{dtr - 1})").font = bold
        for col in range(1, 9):
            ws3.cell(dtr, col).border = top
        for r in range(2, dtr + 1):
            for col in (4, 5, 6):
                ws3.cell(r, col).number_format = "#,##0"
            for col in (7, 8):
                ws3.cell(r, col).number_format = money
        for i, w in enumerate([34, 14, 26, 10, 15, 16, 14, 14], 1):
            ws3.column_dimensions[get_column_letter(i)].width = w

    if console_df is not None and len(console_df):
        ws4 = wb.create_sheet("Console API Usage")
        ws4["A1"] = f"Console/API (platform.claude.com) usage — {console_period}"
        ws4["A1"].font = Font(name=A, bold=True, size=13)
        ws4["A2"] = ("Separate billing relationship from Team seats (Admin API "
                     "cost report) — not blended into Per-User Cost.")
        ws4["A2"].font = base
        hdr_row = 4
        c_headers = ["Workspace", "Model", "Cost Type", "Service Tier",
                     "Amount (USD)"]
        for col, h in enumerate(c_headers, 1):
            c = ws4.cell(hdr_row, col, h)
            c.font, c.fill = hdr_font, hdr_fill
            c.alignment = Alignment(horizontal="center", vertical="center",
                                    wrap_text=True)
        ws4.freeze_panes = f"A{hdr_row + 1}"
        for i, row in console_df.reset_index(drop=True).iterrows():
            r = hdr_row + 1 + i
            vals = [row["workspace"], row["model"], row["cost_type"],
                    row["service_tier"], row["amount_usd"]]
            for col, v in enumerate(vals, 1):
                c = ws4.cell(r, col, v)
                c.font, c.border = base, thin
                if i % 2:
                    c.fill = gray
        ctr = hdr_row + 1 + len(console_df)
        ws4.cell(ctr, 1, "TOTAL").font = bold
        AL = get_column_letter(5)
        ws4.cell(ctr, 5,
                 f"=SUM({AL}{hdr_row + 1}:{AL}{ctr - 1})").font = bold
        for col in range(1, 6):
            ws4.cell(ctr, col).border = top
        for r in range(hdr_row + 1, ctr + 1):
            ws4.cell(r, 5).number_format = money
        for i, w in enumerate([30, 24, 16, 14, 16], 1):
            ws4.column_dimensions[get_column_letter(i)].width = w

    wb.save(out_path)


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("csv", type=Path)
    ap.add_argument("--roster", type=Path, default=None)
    ap.add_argument("--seat-price", type=float, default=30.0)
    ap.add_argument("--premium-price", type=float, default=150.0)
    ap.add_argument("--out", type=Path, default=None)
    ap.add_argument("--no-xlsx", action="store_true")
    ap.add_argument("--console-usage", action="store_true",
                    help="Also fetch platform.claude.com (Console/API) cost "
                         "data via the Admin API (needs ANTHROPIC_ADMIN_KEY)")
    ap.add_argument("--console-start", type=str, default=None,
                    help="Start date YYYY-MM-DD for Console/API data "
                         "(default: same period as the Team CSV filename)")
    ap.add_argument("--console-end", type=str, default=None,
                    help="End date YYYY-MM-DD for Console/API data "
                         "(default: same period as the Team CSV filename)")
    args = ap.parse_args()

    df = pd.read_csv(args.csv)
    fmt = detect_format(df)
    roster = load_roster(args.roster) if args.roster else None
    users, detail = normalize(df, fmt, roster)

    users["seat_fee"] = users["seat_tier"].map(
        lambda t: seat_fee(t, args.seat_price, args.premium_price))
    users["total_cost"] = users["usage_spend"] + users["seat_fee"]
    users = users.sort_values(["total_cost", "usage_spend"],
                              ascending=False).reset_index(drop=True)

    period = period_from_filename(args.csv)
    seats = users["seat_tier"].value_counts().to_dict()
    total_fee, total_spend = users["seat_fee"].sum(), users["usage_spend"].sum()

    print(f"\nClaude Team cost report — {period}  [format: {fmt}"
          f"{', roster joined' if roster is not None else ''}]")
    print(f"Members: {len(users)}  Seats: {seats}")
    print(f"Seat fees: ${total_fee:,.2f}   Usage spend: ${total_spend:,.2f}   "
          f"TOTAL: ${total_fee + total_spend:,.2f}\n")
    cols = ["email", "seat_tier", "usage_spend", "seat_fee", "total_cost"]
    print(users[cols].to_string(
        index=False, formatters={c: "${:,.2f}".format
                                 for c in ("usage_spend", "seat_fee",
                                           "total_cost")}))

    if detail is not None:
        top = detail[detail["net_spend"] > 0].head(5)
        if len(top):
            print("\nTop spend lines (user / product / model):")
            for _, t in top.iterrows():
                print(f"  ${t['net_spend']:>7,.2f}  {t['user_email']}  "
                      f"{t['product']} / {t['model']}")

    if fmt == "spend" and roster is None:
        print("\nNote: spend reports only list users with usage; pass "
              "--roster <members-analytics.csv> to include zero-usage seats "
              "and real seat tiers.")

    console_df, console_period = None, None
    if args.console_usage:
        admin_key = os.environ.get("ANTHROPIC_ADMIN_KEY")
        if not admin_key:
            sys.exit("--console-usage requires the ANTHROPIC_ADMIN_KEY "
                     "environment variable (an Admin API key, "
                     "sk-ant-admin01-..., from platform.claude.com > "
                     "Settings > Admin API keys).")
        c_start = args.console_start
        c_end = args.console_end
        if not c_start or not c_end:
            f_start, f_end = parse_period_dates(args.csv)
            c_start, c_end = c_start or f_start, c_end or f_end
        if not c_start or not c_end:
            sys.exit("Could not determine a date range for --console-usage "
                     "(the CSV filename has none); pass --console-start "
                     "and --console-end explicitly.")
        console_period = f"{c_start} to {c_end}"
        print(f"\nFetching Console/API usage from platform.claude.com for "
              f"{console_period} ...")
        rows = fetch_console_cost_report(admin_key, c_start, c_end)
        ws_names = fetch_console_workspace_names(admin_key)
        console_df = build_console_df(rows, ws_names)

        total_console = console_df["amount_usd"].sum() if len(console_df) else 0.0
        print(f"\nConsole/API (platform.claude.com) usage — {console_period}")
        print(f"Total Console/API spend: ${total_console:,.2f}  "
              "(separate billing relationship from Team seats above — not "
              "included in the TOTAL there)")
        if len(console_df):
            print("Top spend lines (workspace / model / cost type):")
            for _, t in console_df.head(5).iterrows():
                print(f"  ${t['amount_usd']:>7,.2f}  {t['workspace']}  "
                      f"{t['model']} / {t['cost_type']}")

    if not args.no_xlsx:
        out = args.out or args.csv.with_name(
            f"claude-team-per-user-cost_{period.replace(' ', '_')}.xlsx")
        build_xlsx(users, detail, out, args.seat_price, args.premium_price,
                   period, fmt, roster is not None,
                   console_df=console_df, console_period=console_period)
        print(f"\nWrote {out}")
        print("Note: open once in Excel/LibreOffice so formulas calculate "
              "(values are formula-driven, not cached).")


if __name__ == "__main__":
    main()
